"""Apply a restrained competition-report layout to an existing semantic DOCX."""

from argparse import ArgumentParser
from pathlib import Path
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(82, 82, 82)


def set_font(style, name, size, *, bold=False, color=BLACK):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is not None:
        fonts.set(qn("w:eastAsia"), name)
    color_node = rpr.find(qn("w:color"))
    if color_node is not None:
        color_node.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
        for attr in ("themeColor", "themeTint", "themeShade"):
            color_node.attrib.pop(qn(f"w:{attr}"), None)


def remove_border(style):
    ppr = style._element.pPr
    if ppr is not None:
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)


def force_run_color(paragraph, color):
    for run in paragraph.runs:
        run.font.color.rgb = color
        rpr = run._element.get_or_add_rPr()
        node = rpr.find(qn("w:color"))
        if node is not None:
            for attr in ("themeColor", "themeTint", "themeShade"):
                node.attrib.pop(qn(f"w:{attr}"), None)


def set_cell_fill(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shade = tcpr.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        tcpr.append(shade)
    shade.set(qn("w:fill"), color)


def set_cell_margins(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), "110" if side in ("top", "bottom") else "145")
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")


def add_evidence_table(doc, data):
    rows = data["rows"]
    if not rows or any(len(row) != 3 for row in rows):
        raise ValueError("Evidence JSON requires non-empty rows of exactly three strings")
    if len(data.get("headers", ["阶段", "本人行动", "结果与证明"])) != 3:
        raise ValueError("Evidence table needs exactly three headers")
    caption_text = data.get("caption", "阶段证据对照")
    if any(p.text.strip() == caption_text for p in doc.paragraphs):
        raise ValueError("Evidence table caption already exists in the input document")
    target = next((p for p in doc.paragraphs if p.text.strip() == data["before"]), None)
    if target is None:
        raise ValueError(f"Evidence table target heading not found: {data['before']}")

    caption = doc.add_paragraph(caption_text)
    caption.paragraph_format.space_before = Pt(9)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True
    for run in caption.runs:
        run.font.name = "黑体"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = BLACK

    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.style = "Table Grid"
    widths = (Cm(2.4), Cm(6.1), Cm(8.0))
    for row_index, values in enumerate([data.get("headers", ["阶段", "本人行动", "结果与证明"]), *rows]):
        row = table.rows[0] if row_index == 0 else table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            set_cell_margins(cell)
            set_cell_borders(cell)
            set_cell_fill(cell, "353B3B" if row_index == 0 else ("F5F6F5" if row_index % 2 == 0 else "FFFFFF"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.15
                for run in paragraph.runs:
                    run.font.name = "宋体"
                    run.font.size = Pt(9)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else BLACK

    target._p.addprevious(caption._p)
    target._p.addprevious(table._tbl)


def format_report(input_path, output_path, break_before=None, title_text=None, evidence_data=None):
    doc = Document(input_path)
    if title_text:
        title_paragraph = next((p for p in doc.paragraphs if p.style.name == "Title"), None)
        if title_paragraph is None:
            raise ValueError("Input document has no Title paragraph")
        title_paragraph.text = title_text
    normal = doc.styles["Normal"]
    set_font(normal, "宋体", 10.5)
    normal.paragraph_format.line_spacing = 1.27
    normal.paragraph_format.space_after = Pt(4)

    title = doc.styles["Title"]
    set_font(title, "黑体", 16, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(9)
    remove_border(title)

    major = doc.styles["Heading 1"]
    set_font(major, "黑体", 13, bold=True)
    major.paragraph_format.space_before = Pt(11)
    major.paragraph_format.space_after = Pt(5)
    remove_border(major)

    minor = doc.styles["Heading 2"]
    set_font(minor, "黑体", 10.5, bold=True)
    minor.paragraph_format.space_before = Pt(7)
    minor.paragraph_format.space_after = Pt(2)
    remove_border(minor)

    if "Simulation Note" in doc.styles:
        note = doc.styles["Simulation Note"]
        set_font(note, "宋体", 8.5, color=GRAY)
        note.paragraph_format.space_after = Pt(9)

    for paragraph in doc.paragraphs:
        name = paragraph.style.name
        if name == "Title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.keep_with_next = True
            force_run_color(paragraph, BLACK)
        elif name in ("Heading 1", "Heading 2"):
            paragraph.paragraph_format.keep_with_next = True
            force_run_color(paragraph, BLACK)
            if break_before and paragraph.text.strip() == break_before:
                paragraph.paragraph_format.page_break_before = True
        elif name == "Simulation Note":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            force_run_color(paragraph, GRAY)
        else:
            paragraph.paragraph_format.first_line_indent = Cm(0.7)
            paragraph.paragraph_format.widow_control = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for section in doc.sections:
        section.top_margin, section.bottom_margin = Cm(1.85), Cm(1.75)
        section.left_margin, section.right_margin = Cm(2.25), Cm(2.25)
        footer = section.footer.paragraphs[0]
        if "第" in footer.text and "页" in footer.text:
            footer.text = "成长赛道生涯发展报告"
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.name = "宋体"
            run.font.size = Pt(8.5)
            run.font.color.rgb = GRAY

    if evidence_data:
        add_evidence_table(doc, evidence_data)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--page-break-before", dest="break_before")
    parser.add_argument("--title", dest="title_text")
    parser.add_argument("--evidence-json", type=Path)
    arguments = parser.parse_args()
    evidence = json.loads(arguments.evidence_json.read_text(encoding="utf-8")) if arguments.evidence_json else None
    print(format_report(arguments.input, arguments.output, arguments.break_before, arguments.title_text, evidence))
